import unicodedata
from io import BytesIO

from django.conf import settings
from django.contrib.auth.models import User
from django.core.exceptions import ValidationError as ErrorDjango
from django.core.mail import send_mail
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .auth import EsAdministrador
from .carga_masiva import COLUMNAS
from .models import Matricula, Perfil, Seccion
from .usuarios import generar_clave_temporal

TAMANO_MAXIMO = 2 * 1024 * 1024
FILAS_MAXIMAS = 500
LIMITE_DE_LECTURA = 2000


class ErrorArchivo(Exception):
    pass


def _normalizar(texto):
    sin_tildes = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore')
    return sin_tildes.decode().strip().lower()


def _texto(valor):
    return str(valor).strip() if valor is not None else ''


ENCABEZADOS = {_normalizar(c): c for c in COLUMNAS}


def leer_archivo(archivo):
    try:
        libro = load_workbook(BytesIO(archivo.read()), read_only=True, data_only=True)
        hoja = libro['Alumnos'] if 'Alumnos' in libro.sheetnames else libro.worksheets[0]
        todas = list(hoja.iter_rows(values_only=True, max_row=LIMITE_DE_LECTURA))
    except Exception:
        raise ErrorArchivo('No se pudo leer el archivo. Verifica que sea un Excel (.xlsx) válido.')
    if not todas:
        raise ErrorArchivo('El archivo está vacío.')
    indices = {_normalizar(t): i for i, t in enumerate(todas[0]) if t is not None}
    faltantes = [nombre for clave, nombre in ENCABEZADOS.items() if clave not in indices]
    if faltantes:
        raise ErrorArchivo('Faltan las columnas: ' + ', '.join(faltantes) + '.')
    filas = [
        (numero, fila)
        for numero, fila in enumerate(todas[1:], start=2)
        if any(celda not in (None, '') for celda in fila)
    ]
    if len(filas) > FILAS_MAXIMAS:
        raise ErrorArchivo(f'El archivo supera el máximo de {FILAS_MAXIMAS} alumnos por carga.')
    return {clave: indices[clave] for clave in ENCABEZADOS}, filas


def validar_fila(v, vistos):
    errores = []
    if not v['nombres']:
        errores.append('Falta el nombre.')
    if not v['apellidos']:
        errores.append('Faltan los apellidos.')
    correo = v['correo'].lower()
    if not correo:
        errores.append('Falta el correo.')
    else:
        try:
            validate_email(correo)
        except ErrorDjango:
            errores.append('El correo no es válido.')
        else:
            if correo in vistos:
                errores.append('El correo está repetido en el archivo.')
            elif User.objects.filter(username__iexact=correo).exists():
                errores.append('Ya existe un usuario con este correo.')
    seccion = None
    if not v['curso'] or not v['seccion']:
        errores.append('Faltan el curso o la sección.')
    else:
        candidatas = list(
            Seccion.objects.filter(
                activa=True,
                curso__activo=True,
                curso__nombre__iexact=v['curso'],
                nombre__iexact=v['seccion'],
            ).select_related('curso')
        )
        if not candidatas:
            errores.append(
                f'No existe una sección activa "{v["seccion"]}" del curso "{v["curso"]}".'
            )
        elif len(candidatas) > 1:
            errores.append('Hay varias secciones con ese curso y nombre; no se puede elegir una.')
        else:
            seccion = candidatas[0]
            usados = Matricula.objects.filter(seccion=seccion, activa=True).count()
            if usados >= seccion.cupo:
                errores.append('La sección no tiene cupo disponible.')
    return errores, seccion


def crear_alumno(v, seccion):
    correo = v['correo'].lower()
    temporal = generar_clave_temporal()
    with transaction.atomic():
        user = User.objects.create_user(
            username=correo,
            email=correo,
            password=temporal,
            first_name=v['nombres'],
            last_name=v['apellidos'],
        )
        perfil = Perfil.objects.create(user=user, rol=Perfil.Rol.ALUMNO, debe_cambiar_clave=True)
        Matricula.objects.create(alumno=perfil, seccion=seccion)
    return user, temporal


def enviar_credenciales(user, temporal, seccion):
    send_mail(
        'Bienvenido a Umbral',
        (
            f'Hola {user.first_name},\n\n'
            f'Quedaste matriculado en {seccion.curso.nombre} (sección {seccion.nombre}).\n'
            f'Usuario: {user.email}\n'
            f'Contraseña temporal: {temporal}\n\n'
            'Al ingresar se te pedirá crear una contraseña nueva.\n'
        ),
        settings.DEFAULT_FROM_EMAIL,
        [user.email],
        fail_silently=False,
    )


class CargarAlumnosView(APIView):
    permission_classes = [EsAdministrador]
    parser_classes = [MultiPartParser]

    def post(self, request):
        archivo = request.FILES.get('archivo')
        if archivo is None:
            return Response(
                {'archivo': ['Adjunta el archivo Excel.']}, status=status.HTTP_400_BAD_REQUEST
            )
        if not archivo.name.lower().endswith('.xlsx'):
            return Response(
                {'archivo': ['El archivo debe ser un Excel (.xlsx).']},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if archivo.size > TAMANO_MAXIMO:
            return Response(
                {'archivo': ['El archivo no puede superar los 2 MB.']},
                status=status.HTTP_400_BAD_REQUEST,
            )
        try:
            indices, filas = leer_archivo(archivo)
        except ErrorArchivo as error:
            return Response({'archivo': [str(error)]}, status=status.HTTP_400_BAD_REQUEST)

        vistos = set()
        creados, errores, sin_correo = 0, [], []
        for numero, fila in filas:
            v = {
                clave: _texto(fila[i]) if i < len(fila) else ''
                for clave, i in indices.items()
            }
            problemas, seccion = validar_fila(v, vistos)
            if problemas:
                errores.append({'fila': numero, 'correo': v['correo'], 'errores': problemas})
                continue
            user, temporal = crear_alumno(v, seccion)
            vistos.add(v['correo'].lower())
            creados += 1
            try:
                enviar_credenciales(user, temporal, seccion)
            except Exception:
                sin_correo.append(user.email)
        return Response(
            {
                'total': len(filas),
                'creados': creados,
                'con_errores': len(errores),
                'errores': errores,
                'sin_correo': sin_correo,
            }
        )
