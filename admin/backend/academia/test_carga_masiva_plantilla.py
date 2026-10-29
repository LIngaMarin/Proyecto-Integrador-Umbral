from io import BytesIO

from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from .carga_masiva import COLUMNAS
from .models import Perfil
from .test_usuarios import crear_usuario


class PlantillaAlumnosTests(APITestCase):
    def setUp(self):
        self.admin = crear_usuario('admin@test.pe', Perfil.Rol.ADMIN)
        self.docente = crear_usuario('docente@test.pe', Perfil.Rol.DOCENTE)

    def test_admin_descarga_la_plantilla_con_las_columnas(self):
        self.client.force_authenticate(self.admin)
        respuesta = self.client.get('/api/alumnos-masivo/plantilla/')
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual([c.value for c in libro['Alumnos'][1]], COLUMNAS)
        self.assertIn('Instrucciones', libro.sheetnames)

    def test_docente_no_puede_descargar_la_plantilla(self):
        self.client.force_authenticate(self.docente)
        respuesta = self.client.get('/api/alumnos-masivo/plantilla/')
        self.assertEqual(respuesta.status_code, status.HTTP_403_FORBIDDEN)

    def test_sin_sesion_se_rechaza(self):
        respuesta = self.client.get('/api/alumnos-masivo/plantilla/')
        self.assertIn(
            respuesta.status_code,
            (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN),
        )
